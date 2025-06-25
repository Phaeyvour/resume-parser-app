#IMPORT THE NECESSARY LIBRARIES
import streamlit as st
import os
import pdfplumber
import docx
import re
from datetime import datetime
from transformers import pipeline
from fpdf import FPDF

# TEXT EXTRACTION

def extract_text_from_pdf(file_path):
    text = ''
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            text += page.extract_text() or ''
    return text.strip()

def extract_text_from_docx(file_path):
    doc = docx.Document(file_path)
    return '\n'.join([para.text for para in doc.paragraphs]).strip()

def extract_text(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    
    if ext == '.pdf':
        text = extract_text_from_pdf(file_path)
    elif ext == '.docx':
        text = extract_text_from_docx(file_path)
    elif ext == '.txt':
        with open(file_path, 'r', encoding='utf-8') as f:
            text = f.read()
    else:
        raise ValueError("Unsupported file format. Use PDF, DOCX or TXT.")

    # Preprocess to improve link extraction
    text = text.replace('\n', ' ')
    
    return text.strip()
# SUMMARIZATION

@st.cache_resource
def load_summarizer():
    return pipeline("summarization", model="facebook/bart-large-cnn")

summarizer = load_summarizer()

def summarize_with_bart(text):
    text = text[:1024]
    summary = summarizer(text, max_length=130, min_length=30, do_sample=False)
    return summary[0]['summary_text']

# PERSONAL INFO EXTRACTION

def extract_personal_info(text):
    email = re.search(r'[\w\.-]+@[\w\.-]+\.\w+', text)
    phone = re.search(r'(?<!\d)(\+\d{1,3}[\s-]?)?(\(?\d{3,4}\)?[\s.-]?)?\d{3,4}[\s.-]?\d{3,4}(?!\d)', text)
    dob = re.search(r'(\d{1,2}[/-]\d{1,2}[/-]\d{2,4})', text)
    exp = re.search(r'(\d{1,2})\s+(?:years|yrs)[\s\w]*experience', text, re.IGNORECASE)

    # LinkedIn: supports both full and partial formats
    linkedin = re.search(r'(https?://)?(www\.)?linkedin\.com/[^\s\n]+', text, re.IGNORECASE)
    github = re.search(r'(https?://)?(www\.)?github\.com/[^\s\n]+', text, re.IGNORECASE)

    email = email.group(0) if email else "Not found"
    phone = phone.group(0) if phone else "Not found"
    dob_str = dob.group(0) if dob else "Not found"
    linkedin = linkedin.group(0) if linkedin else "Not found"
    github = github.group(0) if github else "Not found"

    # Age calculation
    age = "Not available"
    if dob_str != "Not found":
        for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y", "%d-%m-%y"):
            try:
                dob_dt = datetime.strptime(dob_str, fmt)
                today = datetime.today()
                age = today.year - dob_dt.year - ((today.month, today.day) < (dob_dt.month, dob_dt.day))
                break
            except:
                continue

    experience = f"{exp.group(1)} years" if exp else "Not found"
    return email, phone, dob_str, age, experience, linkedin, github

# SKILLS EXTRACTION & CATEGORIZATION

def extract_skills(text):
    skill_keywords = {
        "Programming": ["python", "java", "c++", "javascript", "html", "css", "sql"],
        "Tools": ["excel", "power bi", "tableau", "git", "docker", "aws"],
        "Soft Skills": ["communication", "leadership", "teamwork", "problem solving"]
    }
    found = {"Programming": [], "Tools": [], "Soft Skills": []}
    text_lower = text.lower()
    for cat, keywords in skill_keywords.items():
        for kw in keywords:
            if kw in text_lower:
                found[cat].append(kw)
    return found

# EXPORT TO PDF

def export_to_pdf(summary, personal_info, skills):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.multi_cell(0, 10, "Resume Summary\n\n" + summary)

    pdf.ln()
    pdf.set_font("Arial", 'B', size=12)
    pdf.cell(0, 10, "\nExtracted Info:", ln=True)
    pdf.set_font("Arial", size=12)
    for key, value in personal_info.items():
        pdf.cell(0, 10, f"{key}: {value}", ln=True)

    pdf.ln()
    pdf.set_font("Arial", 'B', size=12)
    pdf.cell(0, 10, "\nSkills:", ln=True)
    pdf.set_font("Arial", size=12)
    for cat, items in skills.items():
        pdf.cell(0, 10, f"{cat}: {', '.join(items)}", ln=True)

    out_path = "resume_output.pdf"
    pdf.output(out_path)
    return out_path

# PAGE CONFIG

st.set_page_config(page_title="Resume Parser", layout="centered")

# SIDEBAR

st.sidebar.title("⚙️ Settings")
st.sidebar.markdown("""
<div style='text-align: center; font-size: 14px;'>
    This app extracts resume data and summarizes key content.<br>
    Built by <span style='color:gold; font-weight:bold;'>Uzor Favour</span><br>
    using Streamlit.
</div>
""", unsafe_allow_html=True)

# MAIN UI

st.title("📄 Smart Resume Extractor")
st.markdown("Upload your resume and get a quick summary with extracted text.")

# FILE UPLOAD

st.markdown("### 📝 Step 1: Upload Your Resume")
uploaded_file = st.file_uploader("Choose a PDF, DOCX, or TXT file", type=["pdf", "docx", "txt"])

if uploaded_file:
    st.success("✅ Resume uploaded successfully!")

    file_path = os.path.join(os.getcwd(), uploaded_file.name)
    with open(file_path, "wb") as f:
        f.write(uploaded_file.read())

    try:
        with st.spinner("🔍 Hold on a minute my oga..."):
            resume_text = extract_text(file_path)
            summary = summarize_with_bart(resume_text)
            email, phone, dob, age, experience, linkedin, github = extract_personal_info(resume_text)
            skills = extract_skills(resume_text)

        st.success("✅ Extraction complete!")

        st.subheader("📌 Resume Summary")
        st.write(summary)

        st.subheader("🧾 Extracted Info")
        col1, col2 = st.columns(2)
        with col1:
            st.markdown(f"**📧 Email:** {email}")
            st.markdown(f"**📞 Phone:** {phone}")
            st.markdown(f"**🎂 Date of Birth:** {dob}")
            st.markdown(f"**🔗 LinkedIn:** {linkedin}")
        with col2:
            st.markdown(f"**🎈 Age:** {age}")
            st.markdown(f"**💼 Experience:** {experience}")
            st.markdown(f"**🐱 GitHub:** {github}")

        st.subheader("🛠 Skills by Category")
        selected_category = st.selectbox("Select Skill Category", ["All"] + list(skills.keys()))
        if selected_category == "All":
            for cat, items in skills.items():
                st.markdown(f"**{cat}**: {', '.join(items) if items else 'None'}")
        else:
            st.markdown(f"**{selected_category}**: {', '.join(skills[selected_category]) if skills[selected_category] else 'None'}")

        pdf_path = export_to_pdf(summary, {
            "Email": email,
            "Phone": phone,
            "DOB": dob,
            "Age": age,
            "Experience": experience,
            "LinkedIn": linkedin,
            "GitHub": github
        }, skills)

        with open(pdf_path, "rb") as f:
            st.download_button("📥 Download PDF Report", f, file_name="resume_summary.pdf", mime="application/pdf")

        with st.expander("📄 View Full Extracted Text"):
            st.text(resume_text)

        st.balloons()

    except Exception as e:
        st.error(f"❌ Error: {e}")

# STYLE

st.markdown("""
    <style>
    .stButton>button {
        background-color: #4CAF50;
        color: white;
        font-weight: bold;
        border-radius: 12px;
        padding: 0.5em 1em;
    }
    </style>
""", unsafe_allow_html=True)
